#!/usr/bin/env python
# -*- coding: utf-8 -*-
'''
@File    :   file_converter.py
@Author  :   wzm
@Desc    :   文档解析模块：Word转pdf
使用LibreOffice实现word到pdf的转换
LibreOffice 是一个免费、开源、跨平台的办公软件套件，是 Microsoft Office 的免费开源替代品
其跨平台(win/linux/macos)、无 Word 依赖
需要在服务器主机上额外安装LibreOffice软件，然后在python中调用cmd命令行执行转换
'''

import subprocess
import os
import tempfile
from dataclasses import dataclass
from typing import Optional, Union
from docx import Document
from loguru import logger


@dataclass
class ConvertResult:
    """
    转换结果数据类
    """
    success: bool
    file_path: str
    file_type: str
    content: Optional[str] = None
    error: Optional[str] = None


class WordConverter:
    """
    支持 Word 文件转换为 PDF
    """

    SUPPORTED_EXTENSIONS = {'.doc', '.docx'}

    def __init__(self, libreoffice_path: Optional[str] = None):
        """
        初始化转换器
        libreoffice_path: LibreOffice 可执行文件路径，如果不指定则使用默认路径
        """
        if libreoffice_path:
            self.libreoffice_path = libreoffice_path
        else:
            self.libreoffice_path = self._get_default_libreoffice_path()

        self.temp_dir = tempfile.gettempdir()

    def _get_default_libreoffice_path(self) -> str:
        """
        获取默认的 LibreOffice 安装路径
        返回LibreOffice 可执行文件路径
        """
        if os.name == "nt":  # Windows
            return r"C:\Program Files\LibreOffice\program\soffice.exe"
        elif os.name == "posix":  # Linux/macOS
            if os.uname()[0] == "Linux":
                return "libreoffice"
            else:  # macOS
                return "/Applications/LibreOffice.app/Contents/MacOS/soffice"
        else:
            raise Exception("不支持的系统")

    def _is_word_file(self, file_path: str) -> bool:
        """
        判断文件是否为 Word 文件
        """
        _, ext = os.path.splitext(file_path)
        return ext.lower() in self.SUPPORTED_EXTENSIONS

    def word_to_pdf(
        self,
        word_path: str,
        pdf_name: Optional[str] = None,
        outdir: Optional[str] = None
    ) -> Optional[str]:
        """
        使用 LibreOffice 将 Word 文件转换为 PDF
        word_path: Word 文件路径
        pdf_name: 要输出 PDF 文件名，不指定则使用原文件名
        outdir: 输出目录，不指定则使用临时目录
        转换成功返回 PDF 文件路径，失败返回 None
        """
        if not os.path.exists(word_path):
            logger.error(f"Word 文件不存在: {word_path}")
            return None

        output_dir = outdir if outdir else self.temp_dir
        os.makedirs(output_dir, exist_ok=True)

        cmd = [
            self.libreoffice_path,
            "--headless",
            "--invisible",
            "--convert-to", "pdf",
            "--outdir", output_dir,
            word_path
        ]

        try:
            result = subprocess.run(
                cmd,
                check=True,
                capture_output=True,
                encoding="utf-8"
            )

            # LibreOffice 输出名和原文件一致
            base_name = os.path.basename(word_path)
            src_pdf = os.path.join(
                output_dir,
                os.path.splitext(base_name)[0] + ".pdf"
            )

            if pdf_name:
                new_pdf_name = os.path.join(output_dir, f"{pdf_name}.pdf")
                os.rename(src_pdf, new_pdf_name)
                logger.info(f"转换完成：{new_pdf_name}")
                return new_pdf_name
            else:
                logger.info(f"转换完成：{src_pdf}")
                return src_pdf

        except subprocess.CalledProcessError as e:
            logger.error(f"转换失败：{e.stderr}")
            return None
        except FileNotFoundError:
            logger.error(f"未找到 LibreOffice，请确认已安装: {self.libreoffice_path}")
            return None
        except Exception as e:
            logger.error(f"转换异常：{e}")
            return None

    def extract_word_text(self, file_path: str, include_table: bool = False) -> Optional[str]:
        """
        提取 Word 文件的文本内容
        file_path: Word 文件路径
        include_table: 是否包含表格内容
        提取的文本内容，失败返回 None
        """
        if not os.path.exists(file_path):
            logger.error(f"Word 文件不存在: {file_path}")
            return None

        try:
            doc = Document(file_path)
            content = []

            # 提取段落
            for para in doc.paragraphs:
                if para.text.strip():
                    content.append(para.text)

            # 提取表格
            if include_table:
                for table in doc.tables:
                    table_rows = []
                    for row in table.rows:
                        row_cells = []
                        for cell in row.cells:
                            row_cells.append(cell.text.strip())
                        if any(row_cells):
                            table_rows.append("\t".join(row_cells))
                    if table_rows:
                        content.append("\n".join(table_rows))

            return '\n\n'.join(content)

        except Exception as e:
            logger.error(f"提取 Word 内容失败：{e}")
            return None

    def convert(self, word_path: str, include_table: bool = False) -> ConvertResult:
        """
        转换 Word 文件优先尝试转换为 PDF，失败则直接提取文本内容
        """
        if not self._is_word_file(word_path):
            return ConvertResult(
                success=False,
                file_path=word_path,
                file_type="unknown",
                error=f"不支持的文件类型，仅支持: {', '.join(self.SUPPORTED_EXTENSIONS)}"
            )

        # 尝试转换为 PDF
        pdf_path = self.word_to_pdf(word_path)

        if pdf_path and os.path.exists(pdf_path):
            return ConvertResult(
                success=True,
                file_path=pdf_path,
                file_type="pdf"
            )

        # PDF 转换失败，提取文本
        logger.info("PDF 转换失败，尝试提取 Word 文本内容")
        text_content = self.extract_word_text(word_path, include_table)

        if text_content:
            return ConvertResult(
                success=True,
                file_path=word_path,
                file_type="text",
                content=text_content
            )

        return ConvertResult(
            success=False,
            file_path=word_path,
            file_type="word",
            error="PDF 转换失败且文本提取失败"
        )


if __name__ == "__main__":
    converter = WordConverter()

    test_file = "test.docx"

    result = converter.convert(test_file, include_table=True)

